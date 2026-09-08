"""md -> ppt -> md の往復（ラウンドトリップ）ゴールデンテスト."""
import base64
import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from mdflow.document import Document  # noqa: E402
from mdflow import pptx_io  # noqa: E402

SAMPLE = (Path(__file__).resolve().parents[1] / "samples" / "login.md").read_text(
    encoding="utf-8"
)

# 1x1 PNG（テスト用ダミー画像）
_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
)


@pytest.fixture()
def png(tmp_path):
    p = tmp_path / "img.png"
    p.write_bytes(_PNG)
    return p


def test_export_import_roundtrip(tmp_path, png):
    doc = Document(SAMPLE)
    payload = doc.build_payload("flow-login", {"role": "admin", "error_count": 0})

    out = tmp_path / "out.pptx"
    pptx_io.export_ppt(payload, png, out, title="ログイン")

    # 生成物は正しい zip(pptx) であること
    assert zipfile.is_zipfile(out)
    with zipfile.ZipFile(out) as z:
        assert "mdflow/payload.txt" in z.namelist()

    res = pptx_io.import_ppt(out)
    assert res.source == "custom_part"
    assert res.hash_ok is True
    assert res.payload.preset == "管理者・正常"
    assert res.payload.conditions == {"role": "admin", "error_count": 0}
    assert res.payload.mermaid == payload.mermaid


def test_notes_fallback_when_custom_part_removed(tmp_path, png):
    doc = Document(SAMPLE)
    payload = doc.build_payload("flow-login", {"role": "admin", "error_count": 0})
    out = tmp_path / "out.pptx"
    pptx_io.export_ppt(payload, png, out)

    # ビジネス側ツールがカスタムパートを落とした状況を再現
    stripped = tmp_path / "stripped.pptx"
    with zipfile.ZipFile(out) as zin, zipfile.ZipFile(stripped, "w") as zout:
        for item in zin.infolist():
            if item.filename == "mdflow/payload.txt":
                continue
            zout.writestr(item, zin.read(item.filename))

    res = pptx_io.import_ppt(stripped)
    assert res.source == "notes"  # 副系にフォールバック
    assert res.hash_ok is True
    assert res.payload.preset == "管理者・正常"


def test_alt_text_carries_marker(tmp_path, png):
    """第三候補（Alt Text）にもマーカーが載っていること."""
    from pptx import Presentation

    doc = Document(SAMPLE)
    payload = doc.build_payload("flow-login", {"role": "admin", "error_count": 0})
    out = tmp_path / "out.pptx"
    pptx_io.export_ppt(payload, png, out)

    prs = Presentation(str(out))
    descrs = [
        s._element.nvPicPr.cNvPr.get("descr")
        for slide in prs.slides
        for s in slide.shapes
        if s.shape_type == 13  # PICTURE
    ]
    assert any(d and d.startswith("MDFLOW:v1:") for d in descrs)


def test_corrupted_custom_part_raises(tmp_path, png):
    """カスタムパートが壊れた base64 のとき decode で検知できる."""
    from mdflow import payload as pmod

    doc = Document(SAMPLE)
    payload = doc.build_payload("flow-login")
    out = tmp_path / "out.pptx"
    pptx_io.export_ppt(payload, png, out)

    corrupt = tmp_path / "corrupt.pptx"
    with zipfile.ZipFile(out) as zin, zipfile.ZipFile(corrupt, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "mdflow/payload.txt":
                data = b"MDFLOW:v1:!!!notbase64!!!"
            zout.writestr(item, data)

    with pytest.raises(Exception):
        pptx_io.import_ppt(corrupt)
