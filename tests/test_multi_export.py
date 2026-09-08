import io
import zipfile

from PIL import Image
from pptx import Presentation

from mdflow import pptx_io
from mdflow.app import export_bundle


def _png(path, color):
    Image.new("RGB", (80, 40), color).save(path)


def test_export_images_creates_one_slide_per_diagram(tmp_path):
    first, second = tmp_path / "a.png", tmp_path / "b.png"
    _png(first, "red")
    _png(second, "blue")
    output = pptx_io.export_images_ppt([("First", first), ("Second", second)], tmp_path / "all.pptx")
    presentation = Presentation(output)
    assert len(presentation.slides) == 2
    assert "First" in presentation.slides[0].notes_slide.notes_text_frame.text


def test_bundle_contains_svg_png_markdown_and_manifest():
    response = export_bundle({
        "md": "# Document",
        "diagrams": [{"name": "flow one", "svg": "<svg/>",
                      "png_dataurl": "data:image/png;base64,aA=="}],
    })
    with zipfile.ZipFile(io.BytesIO(response.body)) as archive:
        assert set(archive.namelist()) == {
            "flow-one.svg", "flow-one.png", "document.md", "manifest.json"
        }
        assert archive.read("document.md") == b"# Document"
