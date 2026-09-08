"""PPTへのスマートエクスポートと完全復元（双方向）.

埋め込みは三重化（堅牢さ優先）:
1. カスタムパート `mdflow/payload.txt`  … PPTX(zip)内の独立ファイル。不可視・大容量・
   画像を差し替えても消えにくい。＝主。
2. スライドのノート                      … PowerPointネイティブで保存が確実。＝副（保険）。
3. 画像のAlt Text                        … 参考（脆いので第三候補）。

復元は 1 → 2 → 3 の順にフォールバックし、mermaid の hash を照合して改変を検知する。
"""
from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.util import Emu, Inches

from . import payload as _payload
from .payload import Payload

_CUSTOM_PART = "mdflow/payload.txt"
_CT_DEFAULT_TXT = '<Default Extension="txt" ContentType="text/plain"/>'


def export_ppt(
    payload: Payload,
    image_path: str | Path,
    out_path: str | Path,
    *,
    title: str = "",
) -> Path:
    """画像を貼った1枚スライドのPPTを生成し、ペイロードを三重に埋め込む."""
    image_path = Path(image_path)
    out_path = Path(out_path)
    marker = _payload.encode(payload)

    prs = Presentation()
    blank = prs.slide_layouts[6]  # 完全な空白レイアウト
    slide = prs.slides.add_slide(blank)

    # 画像をスライド中央付近に配置（アスペクト比維持のため幅のみ指定）
    pic = slide.shapes.add_picture(str(image_path), Emu(0), Emu(0), width=prs.slide_width)
    if pic.height > prs.slide_height:
        pic.height = prs.slide_height
        pic.width = int(pic.height * (pic.width / pic.height))
    pic.left = int((prs.slide_width - pic.width) / 2)
    pic.top = int((prs.slide_height - pic.height) / 2)

    # (3) Alt Text（参考）
    _set_alt_text(pic, marker)

    # (2) ノート（副・保険）
    notes = slide.notes_slide.notes_text_frame
    header = f"[MdFlow] {title or payload.diagram_id or 'diagram'}\n"
    notes.text = header + marker

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)

    # (1) カスタムパート（主）をzipに注入
    _inject_custom_part(buf.getvalue(), marker, out_path)
    return out_path


def _set_alt_text(pic, text: str) -> None:
    """python-pptx が公開しない descr(Alt Text) を直接設定."""
    cNvPr = pic._element.nvPicPr.cNvPr
    cNvPr.set("descr", text)


def _inject_custom_part(pptx_bytes: bytes, marker: str, out_path: Path) -> None:
    """既存の .pptx(zip) に mdflow/payload.txt を追加し、Content_Types に txt を登録."""
    src = zipfile.ZipFile(io.BytesIO(pptx_bytes), "r")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = _ensure_txt_content_type(data)
            dst.writestr(item, data)
        dst.writestr(_CUSTOM_PART, marker.encode("utf-8"))
    src.close()


def _ensure_txt_content_type(ct_xml: bytes) -> bytes:
    text = ct_xml.decode("utf-8")
    if 'Extension="txt"' in text:
        return ct_xml
    text = text.replace("</Types>", _CT_DEFAULT_TXT + "</Types>")
    return text.encode("utf-8")


@dataclass
class ImportResult:
    payload: Payload
    source: str        # "custom_part" | "notes" | "alt_text"
    hash_ok: bool


def import_ppt(path: str | Path) -> ImportResult:
    """PPTからペイロードを抽出（1→2→3 の順）し、hash照合結果を付けて返す."""
    path = Path(path)

    # (1) カスタムパート
    with zipfile.ZipFile(path, "r") as z:
        if _CUSTOM_PART in z.namelist():
            marker = z.read(_CUSTOM_PART).decode("utf-8")
            pl = _payload.decode(marker)
            return ImportResult(pl, "custom_part", pl.hash_ok())

    # (2) ノート / (3) Alt Text
    prs = Presentation(str(path))
    for slide in prs.slides:
        if slide.has_notes_slide:
            note_text = slide.notes_slide.notes_text_frame.text
            if _payload.find_marker(note_text):
                pl = _payload.decode(note_text)
                return ImportResult(pl, "notes", pl.hash_ok())
        for shape in slide.shapes:
            descr = _get_alt_text(shape)
            if descr and _payload.find_marker(descr):
                pl = _payload.decode(descr)
                return ImportResult(pl, "alt_text", pl.hash_ok())

    raise ValueError("このPPTには MdFlow ペイロードが見つかりません")


def _get_alt_text(shape) -> Optional[str]:
    try:
        return shape._element.nvPicPr.cNvPr.get("descr")
    except AttributeError:
        return None
